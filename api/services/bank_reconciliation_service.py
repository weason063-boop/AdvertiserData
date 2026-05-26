from __future__ import annotations

import io
import json
import math
import re
import uuid
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any
from urllib.parse import quote

import pandas as pd
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from api.database import ensure_bank_reconciliation_batches_table
from api.models import BankReconciliationBatch


DATE_ALIASES = (
    "日期",
    "交易日期",
    "交易日",
    "交易时间",
    "记账日期",
    "记账日",
    "账务日期",
    "发生日期",
    "入账日期",
    "入账日",
    "业务日期",
    "date",
    "transactiondate",
    "postingdate",
    "accountingdate",
)
DEBIT_ALIASES = (
    "借方",
    "借方发生额",
    "借方金额",
    "支出",
    "支出金额",
    "付款",
    "付款金额",
    "debit",
    "withdrawal",
    "outflow",
    "paidout",
)
CREDIT_ALIASES = (
    "贷方",
    "贷方发生额",
    "贷方金额",
    "收入",
    "收入金额",
    "收款",
    "收款金额",
    "credit",
    "deposit",
    "inflow",
    "paidin",
)
SUMMARY_ALIASES = (
    "摘要",
    "备注",
    "用途",
    "交易摘要",
    "对方户名",
    "说明",
    "summary",
    "description",
    "memo",
    "remark",
)
REF_ALIASES = (
    "流水号",
    "交易流水号",
    "凭证号",
    "业务编号",
    "单据编号",
    "reference",
    "refno",
    "voucher",
    "transactionid",
)
ACCOUNT_ALIASES = (
    "账户",
    "账号",
    "银行账号",
    "账户名称",
    "account",
    "accountnumber",
)
CURRENCY_ALIASES = (
    "币种",
    "货币",
    "currency",
)

STATUS_LABELS = {
    "matched": "已匹配",
    "bank_unmatched": "银行未匹配",
    "ledger_unmatched": "会计未匹配",
    "duplicate_pending": "重复待确认",
}
DIRECTION_LABELS = {
    "receipt": "收款",
    "payment": "付款",
}


class BankReconciliationService:
    def __init__(self) -> None:
        self._batch_prefix = "BR"

    async def reconcile_uploads(
        self,
        *,
        month: str | None,
        company: str | None,
        bank_file: UploadFile,
        ledger_file: UploadFile,
        actor: str,
        db: Session,
    ) -> dict[str, Any]:
        bank_bytes = await bank_file.read()
        ledger_bytes = await ledger_file.read()
        if not bank_bytes:
            raise HTTPException(status_code=400, detail="银行流水文件为空")
        if not ledger_bytes:
            raise HTTPException(status_code=400, detail="会计银行明细账文件为空")

        return self.reconcile_bytes(
            month=month,
            company=company,
            bank_bytes=bank_bytes,
            bank_filename=bank_file.filename or "银行流水.xlsx",
            ledger_bytes=ledger_bytes,
            ledger_filename=ledger_file.filename or "会计银行明细账.xlsx",
            actor=actor,
            db=db,
        )

    def reconcile_bytes(
        self,
        *,
        month: str | None,
        company: str | None,
        bank_bytes: bytes,
        bank_filename: str,
        ledger_bytes: bytes,
        ledger_filename: str,
        actor: str,
        db: Session,
    ) -> dict[str, Any]:
        ensure_bank_reconciliation_batches_table()
        parsed_bank = self._parse_workbook(bank_bytes, bank_filename, "bank")
        parsed_ledger = self._parse_workbook(ledger_bytes, ledger_filename, "ledger")
        resolved_month = self._resolve_month(month, parsed_bank["entries"] + parsed_ledger["entries"])
        rows, summary = self._match_entries(parsed_bank["entries"], parsed_ledger["entries"])
        summary.update(
            {
                "month": resolved_month,
                "company": str(company or "").strip(),
                "bank_filename": bank_filename,
                "ledger_filename": ledger_filename,
                "total_bank_rows": parsed_bank["source_row_count"],
                "total_ledger_rows": parsed_ledger["source_row_count"],
                "bank_debit_total": self._float_total(parsed_bank["entries"], "debit"),
                "bank_credit_total": self._float_total(parsed_bank["entries"], "credit"),
                "ledger_debit_total": self._float_total(parsed_ledger["entries"], "debit"),
                "ledger_credit_total": self._float_total(parsed_ledger["entries"], "credit"),
            }
        )
        batch = BankReconciliationBatch(
            batch_no=self._new_batch_no(),
            month=resolved_month,
            company=str(company or "").strip(),
            bank_filename=bank_filename,
            ledger_filename=ledger_filename,
            created_by=actor or "system",
            status="completed",
            total_bank_rows=summary["total_bank_rows"],
            total_ledger_rows=summary["total_ledger_rows"],
            matched_count=summary["matched_count"],
            bank_unmatched_count=summary["bank_unmatched_count"],
            ledger_unmatched_count=summary["ledger_unmatched_count"],
            duplicate_count=summary["duplicate_count"],
            bank_debit_total=summary["bank_debit_total"],
            bank_credit_total=summary["bank_credit_total"],
            ledger_debit_total=summary["ledger_debit_total"],
            ledger_credit_total=summary["ledger_credit_total"],
            matched_amount_total=summary["matched_amount_total"],
            difference_amount_total=summary["difference_amount_total"],
            summary_json=json.dumps(summary, ensure_ascii=False),
            result_rows_json=json.dumps(rows, ensure_ascii=False),
        )
        db.add(batch)
        db.commit()
        db.refresh(batch)
        return self._serialize_batch(batch, include_rows=True)

    def list_batches(self, *, limit: int, month: str | None = None, company: str | None = None, db: Session) -> dict[str, Any]:
        ensure_bank_reconciliation_batches_table()
        safe_limit = max(1, min(100, int(limit or 20)))
        query = db.query(BankReconciliationBatch)
        if month and str(month).strip():
            query = query.filter(BankReconciliationBatch.month == str(month).strip())
        if company and str(company).strip():
            query = query.filter(BankReconciliationBatch.company == str(company).strip())
        
        rows = (
            query
            .order_by(BankReconciliationBatch.id.desc())
            .limit(safe_limit)
            .all()
        )
        return {"rows": [self._serialize_batch(row, include_rows=False) for row in rows], "limit": safe_limit}

    def get_batch(self, batch_id: int, db: Session) -> dict[str, Any]:
        batch = self._get_batch_or_404(batch_id, db)
        return self._serialize_batch(batch, include_rows=True)

    def manual_match(
        self,
        *,
        batch_id: int,
        bank_source_id: str,
        ledger_source_id: str,
        actor: str,
        db: Session,
    ) -> dict[str, Any]:
        batch = self._get_batch_or_404(batch_id, db)
        rows = self._loads_json(batch.result_rows_json, [])
        if not isinstance(rows, list):
            rows = []

        bank_source_id = str(bank_source_id or "").strip()
        ledger_source_id = str(ledger_source_id or "").strip()
        if not bank_source_id or not ledger_source_id:
            raise HTTPException(status_code=400, detail="请选择一条银行流水和一条会计明细")

        bank_row = self._find_source_row(rows, "bank", bank_source_id)
        ledger_row = self._find_source_row(rows, "ledger", ledger_source_id)
        if not bank_row or not ledger_row:
            raise HTTPException(status_code=404, detail="选择的明细不存在或已被匹配")

        bank_entry = bank_row.get("bank")
        ledger_entry = ledger_row.get("ledger")
        if not isinstance(bank_entry, dict) or not isinstance(ledger_entry, dict):
            raise HTTPException(status_code=400, detail="选择的明细无效")
        if bank_row.get("status") == "matched" or ledger_row.get("status") == "matched":
            raise HTTPException(status_code=400, detail="已匹配明细不能重复人工匹配")

        bank_amount = self._parse_amount(bank_entry.get("amount"))
        ledger_amount = self._parse_amount(ledger_entry.get("amount"))
        if bank_amount != ledger_amount:
            raise HTTPException(status_code=400, detail="人工匹配要求银行流水和会计明细金额一致")

        bank_side = str(bank_entry.get("side") or "")
        ledger_side = str(ledger_entry.get("side") or "")
        if not self._is_opposite_side(bank_side, ledger_side):
            raise HTTPException(status_code=400, detail="人工匹配要求银行与会计借贷方向相反")

        direction = "receipt" if bank_side == "credit" else "payment"
        matched_row = self._result_row(
            "matched",
            direction,
            str(bank_entry.get("date") or ledger_entry.get("date") or ""),
            float(bank_amount),
            bank=bank_entry,
            ledger=ledger_entry,
        )
        matched_row["manual_match"] = True
        matched_row["manual_matched_by"] = actor or "system"
        matched_row["manual_matched_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        locked_rows = [
            row
            for row in rows
            if row.get("status") == "matched"
            and not self._row_has_source(row, "bank", bank_source_id)
            and not self._row_has_source(row, "ledger", ledger_source_id)
        ]
        remaining_bank_entries: list[dict[str, Any]] = []
        remaining_ledger_entries: list[dict[str, Any]] = []
        for row in rows:
            if row.get("status") == "matched":
                continue
            if self._row_has_source(row, "bank", bank_source_id) or self._row_has_source(row, "ledger", ledger_source_id):
                continue
            bank = row.get("bank")
            ledger = row.get("ledger")
            if isinstance(bank, dict):
                remaining_bank_entries.append(bank)
            if isinstance(ledger, dict):
                remaining_ledger_entries.append(ledger)

        rebuilt_rows, _ = self._match_entries(remaining_bank_entries, remaining_ledger_entries)
        next_rows = [*locked_rows, matched_row, *rebuilt_rows]
        summary = self._rebuild_summary_from_rows(self._loads_json(batch.summary_json, {}), next_rows)

        batch.summary_json = json.dumps(summary, ensure_ascii=False)
        batch.result_rows_json = json.dumps(next_rows, ensure_ascii=False)
        batch.matched_count = summary["matched_count"]
        batch.bank_unmatched_count = summary["bank_unmatched_count"]
        batch.ledger_unmatched_count = summary["ledger_unmatched_count"]
        batch.duplicate_count = summary["duplicate_count"]
        batch.matched_amount_total = summary["matched_amount_total"]
        batch.difference_amount_total = summary["difference_amount_total"]
        batch.updated_at = datetime.now()
        db.commit()
        db.refresh(batch)
        return self._serialize_batch(batch, include_rows=True)

    def build_report(self, batch_id: int, db: Session) -> tuple[io.BytesIO, str]:
        batch = self._get_batch_or_404(batch_id, db)
        payload = self._serialize_batch(batch, include_rows=True)
        rows = payload["rows"]
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "对账汇总"
        self._write_summary_sheet(summary_sheet, payload)
        self._write_result_sheet(workbook.create_sheet("差异明细"), [r for r in rows if r["status"] != "matched"])
        self._write_result_sheet(workbook.create_sheet("已匹配"), [r for r in rows if r["status"] == "matched"])
        self._write_source_sheet(workbook.create_sheet("银行流水"), self._collect_source_entries(rows, "bank"))
        self._write_source_sheet(workbook.create_sheet("会计明细账"), self._collect_source_entries(rows, "ledger"))
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        filename = f"银行对账结果_{payload['month']}_{payload['batch_no']}.xlsx"
        return buffer, filename

    def content_disposition(self, filename: str) -> str:
        encoded = quote(filename)
        return f"attachment; filename=\"{filename}\"; filename*=UTF-8''{encoded}"

    def _parse_workbook(self, file_bytes: bytes, filename: str, source_type: str) -> dict[str, Any]:
        entries: list[dict[str, Any]] = []
        source_row_count = 0
        if self._is_csv(filename):
            return self._parse_csv(file_bytes, filename, source_type)

        try:
            excel = pd.ExcelFile(io.BytesIO(file_bytes))
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"{filename} 无法读取，请确认是 Excel 文件: {exc}") from exc

        detected_any = False
        for sheet_name in excel.sheet_names:
            try:
                preview = pd.read_excel(excel, sheet_name=sheet_name, header=None, nrows=20, dtype=object)
            except Exception:
                continue
            header_index, mapping = self._detect_header(preview)
            if header_index is None:
                continue
            detected_any = True
            df = pd.read_excel(excel, sheet_name=sheet_name, header=header_index, dtype=object)
            if df.empty:
                continue
            mapping = self._map_columns([str(col) for col in df.columns])
            for offset, row in df.iterrows():
                date_text = self._parse_date(row.get(mapping["date"]))
                debit = self._parse_amount(row.get(mapping["debit"]))
                credit = self._parse_amount(row.get(mapping["credit"]))
                if not date_text or (debit == 0 and credit == 0):
                    continue
                source_row_count += 1
                common = {
                    "source": source_type,
                    "sheet": str(sheet_name),
                    "row_number": int(offset) + int(header_index) + 2,
                    "date": date_text,
                    "summary": self._text(row.get(mapping.get("summary"))) if mapping.get("summary") else "",
                    "ref_no": self._text(row.get(mapping.get("ref_no"))) if mapping.get("ref_no") else "",
                    "account": self._text(row.get(mapping.get("account"))) if mapping.get("account") else "",
                    "currency": self._text(row.get(mapping.get("currency"))) if mapping.get("currency") else "",
                }
                if debit != 0:
                    entries.append({**common, "source_id": f"{source_type}-{len(entries) + 1}", "side": "debit", "debit": float(debit), "credit": 0.0, "amount": float(debit)})
                if credit != 0:
                    entries.append({**common, "source_id": f"{source_type}-{len(entries) + 1}", "side": "credit", "debit": 0.0, "credit": float(credit), "amount": float(credit)})

        if not detected_any:
            raise HTTPException(status_code=400, detail=f"{filename} 未识别到日期、借方、贷方表头")
        if not entries:
            raise HTTPException(status_code=400, detail=f"{filename} 未读取到有效发生额")
        return {"entries": entries, "source_row_count": source_row_count}

    def _parse_csv(self, file_bytes: bytes, filename: str, source_type: str) -> dict[str, Any]:
        raw = None
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "gbk", "gb18030"):
            try:
                raw = pd.read_csv(io.BytesIO(file_bytes), header=None, dtype=object, encoding=encoding)
                break
            except Exception as exc:
                last_error = exc
        if raw is None:
            raise HTTPException(status_code=400, detail=f"{filename} 无法读取 CSV: {last_error}")
        header_index, mapping = self._detect_header(raw.head(20))
        if header_index is None:
            raise HTTPException(status_code=400, detail=f"{filename} 未识别到日期、借方、贷方表头")
        df = raw.iloc[header_index + 1:].copy()
        df.columns = [self._text(value) for value in raw.iloc[header_index].tolist()]
        entries: list[dict[str, Any]] = []
        source_row_count = 0
        mapping = self._map_columns([str(col) for col in df.columns])
        for offset, row in df.iterrows():
            date_text = self._parse_date(row.get(mapping["date"]))
            debit = self._parse_amount(row.get(mapping["debit"]))
            credit = self._parse_amount(row.get(mapping["credit"]))
            if not date_text or (debit == 0 and credit == 0):
                continue
            source_row_count += 1
            common = {
                "source": source_type,
                "sheet": "CSV",
                "row_number": int(offset) + 1,
                "date": date_text,
                "summary": self._text(row.get(mapping.get("summary"))) if mapping.get("summary") else "",
                "ref_no": self._text(row.get(mapping.get("ref_no"))) if mapping.get("ref_no") else "",
                "account": self._text(row.get(mapping.get("account"))) if mapping.get("account") else "",
                "currency": self._text(row.get(mapping.get("currency"))) if mapping.get("currency") else "",
            }
            if debit != 0:
                entries.append({**common, "source_id": f"{source_type}-{len(entries) + 1}", "side": "debit", "debit": float(debit), "credit": 0.0, "amount": float(debit)})
            if credit != 0:
                entries.append({**common, "source_id": f"{source_type}-{len(entries) + 1}", "side": "credit", "debit": 0.0, "credit": float(credit), "amount": float(credit)})
        if not entries:
            raise HTTPException(status_code=400, detail=f"{filename} 未读取到有效发生额")
        return {"entries": entries, "source_row_count": source_row_count}

    def _detect_header(self, preview: pd.DataFrame) -> tuple[int | None, dict[str, str]]:
        for idx, row in preview.iterrows():
            mapping = self._map_columns([self._text(value) for value in row.tolist()])
            if mapping.get("date") and mapping.get("debit") and mapping.get("credit"):
                return int(idx), mapping
        return None, {}

    def _map_columns(self, columns: list[str]) -> dict[str, str]:
        mapping: dict[str, str] = {}
        for col in columns:
            normalized = self._normalize_header(col)
            if not normalized:
                continue
            if "date" not in mapping and self._matches_alias(normalized, DATE_ALIASES):
                mapping["date"] = col
            if "debit" not in mapping and self._matches_alias(normalized, DEBIT_ALIASES):
                mapping["debit"] = col
            if "credit" not in mapping and self._matches_alias(normalized, CREDIT_ALIASES):
                mapping["credit"] = col
            if "summary" not in mapping and self._matches_alias(normalized, SUMMARY_ALIASES):
                mapping["summary"] = col
            if "ref_no" not in mapping and self._matches_alias(normalized, REF_ALIASES):
                mapping["ref_no"] = col
            if "account" not in mapping and self._matches_alias(normalized, ACCOUNT_ALIASES):
                mapping["account"] = col
            if "currency" not in mapping and self._matches_alias(normalized, CURRENCY_ALIASES):
                mapping["currency"] = col
        return mapping

    def _matches_alias(self, normalized: str, aliases: tuple[str, ...]) -> bool:
        for alias in aliases:
            alias_norm = self._normalize_header(alias)
            if alias_norm and (normalized == alias_norm or alias_norm in normalized):
                return True
        return False

    def _match_entries(
        self,
        bank_entries: list[dict[str, Any]],
        ledger_entries: list[dict[str, Any]],
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        bank_groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
        ledger_groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)

        for entry in bank_entries:
            direction = "receipt" if entry["side"] == "credit" else "payment"
            bank_groups[(direction, entry["date"], self._amount_key(entry["amount"]))].append(entry)
        for entry in ledger_entries:
            direction = "receipt" if entry["side"] == "debit" else "payment"
            ledger_groups[(direction, entry["date"], self._amount_key(entry["amount"]))].append(entry)

        result_rows: list[dict[str, Any]] = []
        matched_amount_total = Decimal("0.00")
        difference_amount_total = Decimal("0.00")
        for key in sorted(set(bank_groups) | set(ledger_groups), key=lambda item: (item[1], item[0], item[2])):
            direction, date_text, amount_key = key
            amount = float(Decimal(amount_key))
            banks = bank_groups.get(key, [])
            ledgers = ledger_groups.get(key, [])
            if len(banks) == 1 and len(ledgers) == 1:
                matched_amount_total += Decimal(amount_key)
                result_rows.append(self._result_row("matched", direction, date_text, amount, bank=banks[0], ledger=ledgers[0]))
            elif banks and ledgers:
                for entry in banks:
                    difference_amount_total += Decimal(amount_key)
                    result_rows.append(self._result_row("duplicate_pending", direction, date_text, amount, bank=entry, ledger=None))
                for entry in ledgers:
                    difference_amount_total += Decimal(amount_key)
                    result_rows.append(self._result_row("duplicate_pending", direction, date_text, amount, bank=None, ledger=entry))
            elif banks:
                for entry in banks:
                    difference_amount_total += Decimal(amount_key)
                    result_rows.append(self._result_row("bank_unmatched", direction, date_text, amount, bank=entry, ledger=None))
            else:
                for entry in ledgers:
                    difference_amount_total += Decimal(amount_key)
                    result_rows.append(self._result_row("ledger_unmatched", direction, date_text, amount, bank=None, ledger=entry))

        status_counts = Counter(row["status"] for row in result_rows)
        status_amounts: dict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
        for row in result_rows:
            status_amounts[row["status"]] += Decimal(str(row.get("amount") or 0))
        summary = {
            "matched_count": int(status_counts.get("matched", 0)),
            "bank_unmatched_count": int(status_counts.get("bank_unmatched", 0)),
            "ledger_unmatched_count": int(status_counts.get("ledger_unmatched", 0)),
            "duplicate_count": int(status_counts.get("duplicate_pending", 0)),
            "difference_count": int(
                status_counts.get("bank_unmatched", 0)
                + status_counts.get("ledger_unmatched", 0)
                + status_counts.get("duplicate_pending", 0)
            ),
            "matched_amount_total": float(matched_amount_total),
            "difference_amount_total": float(difference_amount_total),
            "bank_unmatched_amount": float(status_amounts["bank_unmatched"]),
            "ledger_unmatched_amount": float(status_amounts["ledger_unmatched"]),
            "duplicate_amount": float(status_amounts["duplicate_pending"]),
        }
        return result_rows, summary

    def _result_row(
        self,
        status: str,
        direction: str,
        date_text: str,
        amount: float,
        *,
        bank: dict[str, Any] | None,
        ledger: dict[str, Any] | None,
    ) -> dict[str, Any]:
        return {
            "status": status,
            "status_label": STATUS_LABELS[status],
            "direction": direction,
            "direction_label": DIRECTION_LABELS[direction],
            "date": date_text,
            "amount": amount,
            "bank": bank,
            "ledger": ledger,
        }

    def _find_source_row(self, rows: list[dict[str, Any]], source: str, source_id: str) -> dict[str, Any] | None:
        for row in rows:
            if self._row_has_source(row, source, source_id):
                return row
        return None

    def _row_has_source(self, row: dict[str, Any], source: str, source_id: str) -> bool:
        entry = row.get(source)
        return isinstance(entry, dict) and str(entry.get("source_id") or "") == source_id

    def _is_opposite_side(self, bank_side: str, ledger_side: str) -> bool:
        return (bank_side == "credit" and ledger_side == "debit") or (
            bank_side == "debit" and ledger_side == "credit"
        )

    def _rebuild_summary_from_rows(self, previous_summary: dict[str, Any], rows: list[dict[str, Any]]) -> dict[str, Any]:
        summary = dict(previous_summary or {})
        status_counts = Counter(row.get("status") for row in rows)
        status_amounts: dict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
        for row in rows:
            status = str(row.get("status") or "")
            status_amounts[status] += Decimal(str(row.get("amount") or 0))

        summary.update(
            {
                "matched_count": int(status_counts.get("matched", 0)),
                "bank_unmatched_count": int(status_counts.get("bank_unmatched", 0)),
                "ledger_unmatched_count": int(status_counts.get("ledger_unmatched", 0)),
                "duplicate_count": int(status_counts.get("duplicate_pending", 0)),
                "difference_count": int(
                    status_counts.get("bank_unmatched", 0)
                    + status_counts.get("ledger_unmatched", 0)
                    + status_counts.get("duplicate_pending", 0)
                ),
                "matched_amount_total": float(status_amounts["matched"]),
                "difference_amount_total": float(
                    status_amounts["bank_unmatched"]
                    + status_amounts["ledger_unmatched"]
                    + status_amounts["duplicate_pending"]
                ),
                "bank_unmatched_amount": float(status_amounts["bank_unmatched"]),
                "ledger_unmatched_amount": float(status_amounts["ledger_unmatched"]),
                "duplicate_amount": float(status_amounts["duplicate_pending"]),
            }
        )
        return summary

    def _resolve_month(self, month: str | None, entries: list[dict[str, Any]]) -> str:
        if month and str(month).strip():
            normalized = str(month).strip()
            if not re.fullmatch(r"20\d{2}-\d{2}", normalized):
                raise HTTPException(status_code=400, detail="月份格式应为 YYYY-MM")
            return normalized
        month_counts = Counter(str(entry.get("date", ""))[:7] for entry in entries if entry.get("date"))
        month_counts.pop("", None)
        if not month_counts:
            raise HTTPException(status_code=400, detail="无法从日期推断对账月份，请手动选择月份")
        return month_counts.most_common(1)[0][0]

    def _serialize_batch(self, batch: BankReconciliationBatch, *, include_rows: bool) -> dict[str, Any]:
        summary = self._loads_json(batch.summary_json, {})
        payload = {
            "id": batch.id,
            "batch_no": batch.batch_no,
            "month": batch.month,
            "company": batch.company,
            "bank_filename": batch.bank_filename,
            "ledger_filename": batch.ledger_filename,
            "created_by": batch.created_by,
            "status": batch.status,
            "created_at": self._format_datetime(batch.created_at),
            "updated_at": self._format_datetime(batch.updated_at),
            "summary": summary,
        }
        if include_rows:
            payload["rows"] = self._loads_json(batch.result_rows_json, [])
        return payload

    def _get_batch_or_404(self, batch_id: int, db: Session) -> BankReconciliationBatch:
        ensure_bank_reconciliation_batches_table()
        batch = db.query(BankReconciliationBatch).filter(BankReconciliationBatch.id == batch_id).first()
        if not batch:
            raise HTTPException(status_code=404, detail="对账批次不存在")
        return batch

    def _write_summary_sheet(self, sheet, payload: dict[str, Any]) -> None:
        summary = payload["summary"]
        rows = [
            ("批次号", payload["batch_no"]),
            ("月份", payload["month"]),
            ("公司主体", payload.get("company") or ""),
            ("银行流水文件", payload.get("bank_filename") or ""),
            ("会计明细账文件", payload.get("ledger_filename") or ""),
            ("创建人", payload.get("created_by") or ""),
            ("创建时间", payload.get("created_at") or ""),
            ("银行记录数", summary.get("total_bank_rows", 0)),
            ("会计记录数", summary.get("total_ledger_rows", 0)),
            ("已匹配", summary.get("matched_count", 0)),
            ("银行未匹配", summary.get("bank_unmatched_count", 0)),
            ("会计未匹配", summary.get("ledger_unmatched_count", 0)),
            ("重复待确认", summary.get("duplicate_count", 0)),
            ("银行借方合计", summary.get("bank_debit_total", 0)),
            ("银行贷方合计", summary.get("bank_credit_total", 0)),
            ("会计借方合计", summary.get("ledger_debit_total", 0)),
            ("会计贷方合计", summary.get("ledger_credit_total", 0)),
            ("已匹配金额", summary.get("matched_amount_total", 0)),
            ("差异金额", summary.get("difference_amount_total", 0)),
            ("银行未匹配金额", summary.get("bank_unmatched_amount", 0)),
            ("会计未匹配金额", summary.get("ledger_unmatched_amount", 0)),
            ("重复待确认金额", summary.get("duplicate_amount", 0)),
        ]
        sheet.append(["项目", "值"])
        for row in rows:
            sheet.append(list(row))
        self._style_sheet(sheet)

    def _write_result_sheet(self, sheet, rows: list[dict[str, Any]]) -> None:
        headers = [
            "状态",
            "方向",
            "日期",
            "金额",
            "银行日期",
            "银行借方",
            "银行贷方",
            "银行摘要",
            "银行流水号",
            "会计日期",
            "会计借方",
            "会计贷方",
            "会计摘要",
            "会计凭证号",
        ]
        sheet.append(headers)
        for row in rows:
            bank = row.get("bank") or {}
            ledger = row.get("ledger") or {}
            sheet.append(
                [
                    row.get("status_label"),
                    row.get("direction_label"),
                    row.get("date"),
                    row.get("amount"),
                    bank.get("date", ""),
                    bank.get("debit", 0),
                    bank.get("credit", 0),
                    bank.get("summary", ""),
                    bank.get("ref_no", ""),
                    ledger.get("date", ""),
                    ledger.get("debit", 0),
                    ledger.get("credit", 0),
                    ledger.get("summary", ""),
                    ledger.get("ref_no", ""),
                ]
            )
        self._style_sheet(sheet)

    def _write_source_sheet(self, sheet, entries: list[dict[str, Any]]) -> None:
        sheet.append(["日期", "借方", "贷方", "摘要", "流水/凭证号", "账户", "币种", "Sheet", "行号"])
        for entry in entries:
            sheet.append(
                [
                    entry.get("date", ""),
                    entry.get("debit", 0),
                    entry.get("credit", 0),
                    entry.get("summary", ""),
                    entry.get("ref_no", ""),
                    entry.get("account", ""),
                    entry.get("currency", ""),
                    entry.get("sheet", ""),
                    entry.get("row_number", ""),
                ]
            )
        self._style_sheet(sheet)

    def _collect_source_entries(self, rows: list[dict[str, Any]], source: str) -> list[dict[str, Any]]:
        seen: set[str] = set()
        entries: list[dict[str, Any]] = []
        for row in rows:
            entry = row.get(source)
            if not isinstance(entry, dict):
                continue
            key = str(entry.get("source_id") or f"{source}-{len(entries)}")
            if key in seen:
                continue
            seen.add(key)
            entries.append(entry)
        return entries

    def _style_sheet(self, sheet) -> None:
        fill = PatternFill("solid", fgColor="E2E8F0")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        for column in sheet.columns:
            max_len = 10
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value) + 2, 32))
            sheet.column_dimensions[column_letter].width = max_len
        sheet.freeze_panes = "A2"

    def _float_total(self, entries: list[dict[str, Any]], side: str) -> float:
        total = Decimal("0.00")
        for entry in entries:
            if entry.get("side") == side:
                total += Decimal(str(entry.get("amount") or 0))
        return float(total)

    def _amount_key(self, value: Any) -> str:
        return str(self._parse_amount(value))

    def _parse_amount(self, value: Any) -> Decimal:
        if value is None:
            return Decimal("0.00")
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return Decimal("0.00")
        text = str(value).strip()
        if not text or text.lower() in {"nan", "none", "-"}:
            return Decimal("0.00")
        negative = text.startswith("(") and text.endswith(")")
        text = text.replace(",", "").replace("￥", "").replace("$", "").replace("¥", "")
        text = text.replace(" ", "").replace("\u3000", "")
        text = text.strip("()")
        try:
            amount = Decimal(text)
        except InvalidOperation:
            match = re.search(r"-?\d+(?:\.\d+)?", text)
            if not match:
                return Decimal("0.00")
            amount = Decimal(match.group(0))
        if negative:
            amount = -amount
        return amount.copy_abs().quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    def _parse_date(self, value: Any) -> str | None:
        if value is None:
            return None
        if isinstance(value, float) and math.isnan(value):
            return None
        if isinstance(value, (int, float)) and 20000 <= float(value) <= 60000:
            parsed_serial = pd.to_datetime(value, unit="D", origin="1899-12-30", errors="coerce")
            if not pd.isna(parsed_serial):
                return parsed_serial.strftime("%Y-%m-%d")
        try:
            parsed = pd.to_datetime(value, errors="coerce")
        except Exception:
            parsed = None
        if parsed is not None and not pd.isna(parsed):
            return parsed.strftime("%Y-%m-%d")

        text = str(value).strip()
        if not text:
            return None
        match = re.search(r"(20\d{2})[年./\-\s]*(\d{1,2})[月./\-\s]*(\d{1,2})", text)
        if match:
            year, month, day = (int(match.group(1)), int(match.group(2)), int(match.group(3)))
            try:
                return datetime(year, month, day).strftime("%Y-%m-%d")
            except ValueError:
                return None
        return None

    def _normalize_header(self, value: Any) -> str:
        return re.sub(r"[\s_（）()【】\[\]:：/\\\-]+", "", self._text(value).lower())

    def _text(self, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and math.isnan(value):
            return ""
        return str(value).strip()

    def _loads_json(self, raw: str | None, fallback: Any) -> Any:
        if not raw:
            return fallback
        try:
            return json.loads(raw)
        except Exception:
            return fallback

    def _format_datetime(self, value: Any) -> str | None:
        if not value:
            return None
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return str(value)

    def _new_batch_no(self) -> str:
        return f"{self._batch_prefix}{datetime.now().strftime('%Y%m%d%H%M%S')}{uuid.uuid4().hex[:6].upper()}"

    def _is_csv(self, filename: str) -> bool:
        return str(filename or "").lower().endswith(".csv")
