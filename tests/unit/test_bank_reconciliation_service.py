from io import BytesIO

from openpyxl import Workbook, load_workbook

from api.models import BankReconciliationBatch
from api.services.bank_reconciliation_service import BankReconciliationService


def _xlsx_bytes(headers, rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append([""])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_bank_reconciliation_matches_reverse_directions_and_flags_differences(db_session):
    bank_bytes = _xlsx_bytes(
        ["交易日期", "借方发生额", "贷方发生额", "摘要", "流水号"],
        [
            ["2026-03-01", "", 1000, "客户回款", "B001"],
            ["2026-03-02", 200, "", "支付媒体款", "B002"],
            ["2026-03-03", "", 500, "银行独有", "B003"],
            ["2026-03-04", "", 100, "重复1", "B004"],
            ["2026-03-04", "", 100, "重复2", "B005"],
        ],
    )
    ledger_bytes = _xlsx_bytes(
        ["记账日期", "借方金额", "贷方金额", "摘要", "凭证号"],
        [
            ["2026-03-01", 1000, "", "客户回款", "L001"],
            ["2026-03-02", "", 200, "支付媒体款", "L002"],
            ["2026-03-03", "", 999, "会计独有", "L003"],
            ["2026-03-04", 100, "", "重复", "L004"],
        ],
    )

    result = BankReconciliationService().reconcile_bytes(
        month="2026-03",
        bank_bytes=bank_bytes,
        bank_filename="bank.xlsx",
        ledger_bytes=ledger_bytes,
        ledger_filename="ledger.xlsx",
        actor="tester",
        db=db_session,
    )

    summary = result["summary"]
    assert summary["matched_count"] == 2
    assert summary["bank_unmatched_count"] == 1
    assert summary["ledger_unmatched_count"] == 1
    assert summary["duplicate_count"] == 3
    assert summary["bank_credit_total"] == 1700.0
    assert summary["bank_debit_total"] == 200.0
    assert summary["ledger_debit_total"] == 1100.0
    assert summary["ledger_credit_total"] == 1199.0
    assert {row["status"] for row in result["rows"]} == {
        "matched",
        "bank_unmatched",
        "ledger_unmatched",
        "duplicate_pending",
    }

    saved = db_session.query(BankReconciliationBatch).one()
    assert saved.batch_no == result["batch_no"]
    assert saved.month == "2026-03"


def test_bank_reconciliation_export_contains_summary_and_detail_sheets(db_session):
    bank_bytes = _xlsx_bytes(
        ["日期", "借方", "贷方", "摘要"],
        [["2026-03-01", "", 1000, "客户回款"]],
    )
    ledger_bytes = _xlsx_bytes(
        ["日期", "借方", "贷方", "摘要"],
        [["2026-03-01", 1000, "", "客户回款"]],
    )
    service = BankReconciliationService()
    result = service.reconcile_bytes(
        month=None,
        bank_bytes=bank_bytes,
        bank_filename="bank.xlsx",
        ledger_bytes=ledger_bytes,
        ledger_filename="ledger.xlsx",
        actor="tester",
        db=db_session,
    )

    buffer, filename = service.build_report(result["id"], db=db_session)
    workbook = load_workbook(buffer)

    assert filename.startswith("银行对账结果_2026-03_")
    assert set(workbook.sheetnames) == {"对账汇总", "差异明细", "已匹配", "银行流水", "会计明细账"}
    assert workbook["对账汇总"]["A1"].value == "项目"
    assert workbook["已匹配"].max_row == 2


def test_bank_reconciliation_reads_cmb_short_trade_day_header(db_session):
    bank_bytes = _xlsx_bytes(
        ["交易日", "交易类型", "借方金额", "贷方金额", "摘要", "用途", "收/付方名称"],
        [
            ["2026/4/1", "提回对公户收款", "", "3,684.00", "信息服务费", "StarLord", "深圳市星洛得科技有限公司"],
            ["2026/4/1", "代发款项", "330,000.00", "", "备用金", "备用金", ""],
        ],
    )
    ledger_bytes = _xlsx_bytes(
        ["记账日期", "借方金额", "贷方金额", "摘要", "凭证号"],
        [
            ["2026/4/1", 3684, "", "信息服务费", "L001"],
            ["2026/4/1", "", 330000, "备用金", "L002"],
        ],
    )

    result = BankReconciliationService().reconcile_bytes(
        month=None,
        bank_bytes=bank_bytes,
        bank_filename="招行基本户流水_2023年4月.xlsx",
        ledger_bytes=ledger_bytes,
        ledger_filename="明细账.xls",
        actor="tester",
        db=db_session,
    )

    assert result["month"] == "2026-04"
    assert result["summary"]["total_bank_rows"] == 2
    assert result["summary"]["matched_count"] == 2
    assert result["summary"]["bank_credit_total"] == 3684.0
    assert result["summary"]["bank_debit_total"] == 330000.0
